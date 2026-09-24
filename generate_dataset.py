"""
generate_dataset.py — builds raw_dataset.xlsx: synthetic-but-realistic data shaped EXACTLY like
the EqualAccess Portal's real Supabase tables (public.pwd_users, public.jobs), plus one proposed
TARGET-label sheet, because the live database currently has no explicit match/like/approval
column (confirmed by reading supabase/migrations/*.sql — see the Data_Dictionary sheet for the
full audit trail of every column's source and whether the recommendation engine reads it).

Sheets written, in order:
  1. pwd_users              — every column of public.pwd_users, in table-definition order (30 rows)
  2. jobs                   — every column of public.jobs, in table-definition order (12 rows)
  3. recommendation_labels  — PROPOSED target label (not an existing table): pwd_user_id, job_id,
                               saved (0/1). Modeled on pwd_users.saved_job_ids (the bookmark button
                               already in the app), simulated probabilistically from skill overlap,
                               disability fit and education fit — not a hard rule — so the label
                               isn't trivially separable. (120 rows)
  4. Data_Dictionary        — column-by-column reference for every sheet above, with source table,
                               data type, allowed values, and whether the current engine uses it.

Sheet/column names for pwd_users and jobs are EXACTLY the database's — nothing added or removed.
No real personal information: all names, contacts and IDs are synthetic.

Usage: python generate_dataset.py
Output: raw_dataset.xlsx (this directory)
"""
import json
import random
from datetime import date, timedelta

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(42)
np.random.seed(42)

N_PWD = 30
N_JOBS = 12
N_LABELS = 120

# ── Catalog values (mirrors src/lib/catalog.ts exactly) ──────────────────────
BARANGAYS = [
    'Anos', 'Bagong Silang', 'Bambang', 'Batong Malake', 'Baybayin', 'Bayog', 'Lalakay',
    'Maahas', 'Malinta', 'Mayondon', 'Putho-Tuntungin', 'San Antonio', 'Tadlac', 'Timugan',
]
DISABILITY_TYPES = [
    'Cancer (RA 11215)', 'Deaf or Hard of Hearing', 'Intellectual Disability', 'Learning Disability',
    'Mental Disability', 'Physical Disability', 'Psychosocial Disability', 'Rare Disease (RA 10747)',
    'Speech & Language Impairment', 'Visual Disability', 'Other',
]
EDUCATION_LEVELS = [
    'Elementary', 'High School Level', 'High School Graduate', 'Senior High School Graduate',
    'Vocational', 'College Level', 'College Graduate', 'Post Graduate',
]
EDUCATION_RANK = {label: i + 1 for i, label in enumerate(EDUCATION_LEVELS)}
EMPLOYMENT_TYPES = ['Full-time', 'Part-time', 'Contractual', 'Freelance']
LEGACY_JOB_TYPES = ['Full-time', 'Part-time', 'Contract', 'Remote']  # old `jobs.type` vocabulary
WORK_ARRANGEMENTS = ['On-site', 'Hybrid', 'Remote']
JOB_STATUSES = ['Draft', 'Open', 'Closed', 'Archived']
VERIFICATION_STATUSES = ['Pending', 'Verified', 'Rejected']
ACCOMMODATIONS = [
    'Wheelchair-accessible workplace', 'Screen-reader-compatible tools', 'Sign-language interpreter',
    'Flexible hours', 'Remote work', 'Quiet workspace', 'Written instructions',
    'Assistive technology provided',
]
# Canonical skill vocabulary (src/lib/recommend/taxonomy.json), Title Case for readability.
SKILL_POOL = [
    'Microsoft Office', 'Computer Literacy', 'Customer Service', 'Communication', 'Data Entry',
    'Typing', 'Attention to Detail', 'Organizing Files', 'Administration', 'Social Media', 'Design',
    'Writing', 'Photography', 'Video Editing', 'English Proficiency', 'Accounting', 'Research',
    'Web Development', 'IT Support', 'Horticulture', 'Manual Labor', 'Carpentry', 'Cooking', 'Baking',
    'Food Safety', 'Food Processing', 'Sales', 'Cashiering', 'Tailoring', 'Teaching', 'Sign Language',
]

FIRST_NAMES = [
    'Maria', 'Juan', 'Ana', 'Jose', 'Liza', 'Carlos', 'Rosa', 'Pedro', 'Carmen', 'Antonio',
    'Elena', 'Ramon', 'Teresa', 'Manuel', 'Grace', 'Ricardo', 'Josefina', 'Danilo', 'Corazon',
    'Ernesto', 'Luz', 'Rodolfo', 'Angelica', 'Ferdinand', 'Precious', 'Benjamin', 'Divina',
    'Marlon', 'Estrella', 'Reynaldo',
]
LAST_NAMES = [
    'Santos', 'Reyes', 'Cruz', 'Bautista', 'Garcia', 'Mendoza', 'Torres', 'Flores', 'Ramos',
    'Gonzales', 'Delos Santos', 'Aquino', 'Villanueva', 'Castillo', 'Rivera', 'Navarro', 'Pascual',
    'Fernandez', 'Salazar', 'Domingo', 'Marquez', 'Ocampo', 'De Guzman', 'Padilla', 'Ignacio',
    'Robles', 'Valdez', 'Serrano', 'Espino', 'Bernardo',
]

JOB_CATALOG = [
    # (title, category, skills, min_education)
    ('Data Entry Clerk', 'Administrative', ['Data Entry', 'Typing', 'Microsoft Office', 'Attention to Detail'], 'Vocational'),
    ('Customer Service Representative', 'Administrative', ['Customer Service', 'Communication', 'English Proficiency'], 'Senior High School Graduate'),
    ('Graphic Designer', 'Creative', ['Design', 'Photography', 'Social Media'], 'College Level'),
    ('Kitchen Assistant', 'Food Service', ['Cooking', 'Food Safety', 'Food Processing'], 'High School Graduate'),
    ('Junior Web Developer', 'Technical', ['Web Development', 'IT Support', 'Computer Literacy'], 'College Graduate'),
    ('Records Assistant', 'Administrative', ['Organizing Files', 'Data Entry', 'Administration'], 'Senior High School Graduate'),
    ('Baker', 'Food Service', ['Baking', 'Food Safety', 'Attention to Detail'], 'High School Graduate'),
    ('Social Media Assistant', 'Creative', ['Social Media', 'Writing', 'Communication'], 'College Level'),
    ('Library Assistant', 'Administrative', ['Organizing Files', 'Research', 'Computer Literacy'], 'Vocational'),
    ('Landscaping Helper', 'Manual', ['Horticulture', 'Manual Labor'], 'Elementary'),
    ('Bookkeeper', 'Administrative', ['Accounting', 'Microsoft Office', 'Attention to Detail'], 'College Graduate'),
    ('Tailoring Assistant', 'Manual', ['Tailoring', 'Attention to Detail'], 'High School Graduate'),
]
COMPANY_NAMES = [
    'Municipal Records Section', 'Los Baños Public Market', 'PDAO Communications Office',
    'Bayog Community Bakeshop', 'UPLB Extension Office', 'Malinta Parish Cooperative',
    'Los Baños Public Library', 'Municipal Accounting Office', 'Batong Malake Web Studio',
    'Anos Agri Supply', 'Mayondon Tailoring Shop', 'Los Baños Health Center',
]


def rand_date(start: date, end: date) -> str:
    delta = (end - start).days
    return (start + timedelta(days=random.randint(0, max(delta, 0)))).isoformat()


def jlist(items) -> str:
    """A jsonb array column, rendered as it is actually stored: a JSON array string."""
    return json.dumps(list(items), ensure_ascii=False)


# ── 1. pwd_users (30 rows) ────────────────────────────────────────────────────
pwd_rows = []
pwd_skill_sets = []  # kept alongside for the label sheet's skill-overlap calculation
for i in range(1, N_PWD + 1):
    pid = f'PWD-LB-2026-{i:05d}'
    first, last = random.choice(FIRST_NAMES), random.choice(LAST_NAMES)
    name = f'{first} {last}'
    barangay = f'Brgy. {random.choice(BARANGAYS)}'
    disability = random.choice(DISABILITY_TYPES)
    skills = random.sample(SKILL_POOL, k=random.randint(3, 6))
    pwd_skill_sets.append(skills)
    education_level = random.choices(EDUCATION_LEVELS, weights=[3, 6, 10, 10, 8, 8, 10, 3])[0]
    verification = random.choices(VERIFICATION_STATUSES, weights=[2, 7, 1])[0]
    pref_job_types = random.sample(EMPLOYMENT_TYPES, k=random.randint(0, 2))
    pref_work_setup = random.sample(WORK_ARRANGEMENTS, k=random.randint(0, 2))
    accommodations = random.sample(ACCOMMODATIONS, k=random.randint(0, 2))
    accessibility_needs = random.sample(ACCOMMODATIONS, k=random.randint(0, 1))

    pwd_rows.append({
        'id': pid,
        'username': f'{first.lower()}.{last.lower().replace(" ", "")}{i:02d}',
        'password': 'pwd123',  # uniform synthetic demo password, matches the portal's own demo accounts
        'name': name,
        'address': f'{random.randint(1, 200)} {random.choice(["Rizal", "Mabini", "Bonifacio", "Burgos", "Luna"])} St.',
        'barangay': barangay,
        'age': random.randint(18, 65),
        'contact': f'+63 9{random.randint(10,99)} {random.randint(100,999)} {random.randint(1000,9999)}',
        'email': f'{first.lower()}.{last.lower().replace(" ", "")}{i:02d}@example.com',
        'disability_type': disability,
        'verification_status': verification,
        'date_registered': rand_date(date(2024, 1, 1), date(2026, 8, 31)),
        'pwd_id_number': f'LB-{disability[:3].upper()}-2026-{i:05d}',
        'avatar': '',
        'active': True if verification != 'Rejected' else False,
        'skills': jlist(skills),
        'education': '' if random.random() < 0.5 else f'{random.choice(["BS", "AB", "BA"])} {random.choice(["Information Technology", "Business Administration", "Education", "Accountancy", "Psychology"])}',
        'work_experience': '' if random.random() < 0.4 else f'{random.randint(0, 5)} year(s) as a {random.choice(["barangay records assistant", "sari-sari store helper", "OJT intern", "part-time cashier"])}',
        'years_of_experience': random.choice([0, 0, 1, 1, 2, 3, 5, None]),
        'certifications': jlist(random.sample(['NC II Computer Servicing', 'NC II Bread and Pastry', 'TESDA Data Entry', 'First Aid Certificate'], k=random.randint(0, 2))),
        'job_interests': jlist(random.sample([t for t, *_ in JOB_CATALOG], k=random.randint(0, 2))),
        'preferred_job_types': jlist(pref_job_types),
        'preferred_work_setup': jlist(pref_work_setup),
        'functional_capabilities': jlist(random.sample(['Computer-based tasks', 'Seated work', 'Verbal communication', 'Manual tasks', 'Standing for short periods'], k=random.randint(0, 3))),
        'accessibility_needs': jlist(accessibility_needs),
        'accommodation_requirements': jlist(accommodations),
        'preferred_location': barangay if random.random() < 0.6 else 'Los Baños, Laguna',
        'education_level': education_level,
        'saved_job_ids': jlist([]),  # filled in after recommendation_labels is generated
        'deleted_at': '',
    })
pwd_df = pd.DataFrame(pwd_rows)

# ── 2. jobs (12 rows) ──────────────────────────────────────────────────────────
job_rows = []
for i, (title, category, skills, min_education) in enumerate(JOB_CATALOG, start=1):
    jid = f'JOB-{i:03d}'
    employment_type = random.choice(EMPLOYMENT_TYPES)
    work_arrangement = random.choices(WORK_ARRANGEMENTS, weights=[6, 2, 2])[0]
    legacy_type = 'Remote' if work_arrangement == 'Remote' else random.choice(LEGACY_JOB_TYPES[:3])
    posted = rand_date(date(2026, 5, 1), date(2026, 8, 1))
    deadline = (date.fromisoformat(posted) + timedelta(days=random.randint(30, 90))).isoformat()
    suitable = [] if random.random() < 0.7 else random.sample(DISABILITY_TYPES[:-1], k=random.randint(1, 2))
    accommodations = random.sample(ACCOMMODATIONS, k=random.randint(0, 3))

    job_rows.append({
        'id': jid,
        'title': title,
        'company': COMPANY_NAMES[i - 1],
        'location': random.choice([f'Brgy. {b}' for b in BARANGAYS] + ['Los Baños, Laguna']),
        'type': legacy_type,
        'category': category,
        'salary': f'₱{random.randint(10, 18)},000 – ₱{random.randint(19, 25)},000',
        'description': f'Responsibilities include tasks typical of a {title.lower()} role at {COMPANY_NAMES[i - 1]}.',
        'skills': jlist(skills),
        'preferred_skills': jlist(random.sample(SKILL_POOL, k=random.randint(0, 2))),
        'education_requirement': min_education,
        'experience_requirement': random.choice(['', 'No experience required', '6 months experience', '1 year experience']),
        'work_setup': work_arrangement,
        'workplace_conditions': jlist(random.sample(['Air-conditioned office', 'Ground floor access', 'Outdoor work'], k=random.randint(0, 1))),
        'accessibility_info': random.choice(['', 'Ground-floor workstation available.', 'Accessible restroom on-site.']),
        'accessibility_features': jlist(random.sample(['Ramp access', 'Elevator', 'Accessible restroom'], k=random.randint(0, 1))),
        'physical_requirements': jlist(random.sample(['Prolonged sitting', 'Light lifting', 'Standing for long periods'], k=random.randint(0, 1))),
        'communication_requirements': jlist(random.sample(['Verbal communication', 'Written communication'], k=random.randint(0, 1))),
        'functional_requirements': jlist([]),
        'accommodation_support': '' if not accommodations else f'Employer can provide: {", ".join(accommodations)}.',
        'posted_date': posted,
        'deadline': deadline,
        'status': random.choices(JOB_STATUSES, weights=[1, 7, 1, 1])[0],
        'match_percent': None,
        'match_reasons': jlist([]),
        'screen_or_visual_demands': jlist(random.sample(['Prolonged screen use', 'Fine print reading'], k=random.randint(0, 1))),
        'employment_type': employment_type,
        'work_arrangement': work_arrangement,
        'min_education': min_education,
        'suitable_disabilities': jlist(suitable),
        'accommodations': jlist(accommodations),
        'slots': random.randint(1, 5),
    })
jobs_df = pd.DataFrame(job_rows)

# ── 3. recommendation_labels (120 rows) — PROPOSED target, not an existing table ──
# Simulated implicit feedback ("saved this job?"), modeled on pwd_users.saved_job_ids.
# Probability of saving blends skill overlap, disability fit and education fit, then a
# Bernoulli draw — so the label is realistic (noisy) rather than a deterministic rule.
all_pairs = [(p, j) for p in range(N_PWD) for j in range(N_JOBS)]
random.shuffle(all_pairs)
chosen_pairs = all_pairs[:N_LABELS]

label_rows = []
saved_by_pwd = {i: [] for i in range(N_PWD)}
for p_idx, j_idx in chosen_pairs:
    pwd = pwd_rows[p_idx]
    job = job_rows[j_idx]
    applicant_skills = set(pwd_skill_sets[p_idx])
    required_skills = set(json.loads(job['skills']))
    skill_overlap = len(applicant_skills & required_skills) / max(1, len(required_skills))

    suitable = json.loads(job['suitable_disabilities'])
    disability_fit = 1.0 if not suitable or pwd['disability_type'] in suitable else 0.3

    applicant_rank = EDUCATION_RANK.get(pwd['education_level'], 1)
    required_rank = EDUCATION_RANK.get(job['min_education'], 1)
    education_fit = 1.0 if applicant_rank >= required_rank else 0.4

    prob = 0.5 * skill_overlap + 0.3 * disability_fit + 0.2 * education_fit
    prob = float(np.clip(prob, 0.05, 0.95))
    saved = int(np.random.binomial(1, prob))

    label_rows.append({'pwd_user_id': pwd['id'], 'job_id': job['id'], 'saved': saved})
    if saved:
        saved_by_pwd[p_idx].append(job['id'])

labels_df = pd.DataFrame(label_rows)

# Keep pwd_users.saved_job_ids consistent with the labels sheet (same foreign keys).
for i, row in enumerate(pwd_rows):
    row['saved_job_ids'] = jlist(saved_by_pwd[i])
pwd_df = pd.DataFrame(pwd_rows)

# ── 4. Data_Dictionary ─────────────────────────────────────────────────────────
def dd_rows_for(df: pd.DataFrame, sheet: str, notes: dict) -> list:
    rows = []
    for col in df.columns:
        sample = df[col].dropna().astype(str)
        sample_val = sample.iloc[0] if len(sample) else ''
        rows.append({
            'Sheet': sheet,
            'Column': col,
            'Data_Type': str(df[col].dtype),
            'Example_Value': sample_val[:60],
            'Notes': notes.get(col, ''),
        })
    return rows


pwd_notes = {
    'id': 'Primary key (public.pwd_users.id).', 'password': 'Synthetic placeholder only — never real credentials.',
    'barangay': 'One of the 14 official Los Baños barangays. Used for the Jobs page location filter, not scored.',
    'disability_type': 'FEATURE — scored (suitability component).',
    'skills': 'FEATURE — scored (skills component). JSON array, free text drawn from the app\'s skill taxonomy.',
    'education_level': 'FEATURE — scored (education component).',
    'accommodation_requirements': 'FEATURE — scored (suitability component), combined with accessibility_needs.',
    'accessibility_needs': 'FEATURE — combined into the suitability score with accommodation_requirements.',
    'preferred_job_types': 'Collected by the profile form; NOT scored by the current engine.',
    'preferred_work_setup': 'Collected by the profile form; NOT scored by the current engine.',
    'years_of_experience': 'Stored only — not read by the recommendation engine.',
    'certifications': 'Stored only — not read by the recommendation engine.',
    'job_interests': 'Stored only — not read by the recommendation engine.',
    'functional_capabilities': 'Stored only — not read by the recommendation engine.',
    'preferred_location': 'Stored only — not read by the recommendation engine (barangay drives the location filter instead).',
    'saved_job_ids': 'Bookmark button on the Jobs page. Basis for the proposed recommendation_labels target below.',
    'deleted_at': 'Soft-delete marker; non-empty excludes the row from every dashboard count.',
}
jobs_notes = {
    'id': 'Primary key (public.jobs.id).',
    'skills': 'FEATURE — required skills, scored (skills component).',
    'min_education': 'FEATURE — scored (education component). Current column (superseded education_requirement).',
    'suitable_disabilities': 'FEATURE — hard filter + suitability boost. Empty = open to all disability types.',
    'accommodations': 'FEATURE — scored (suitability component).',
    'employment_type': 'Collected/displayed; NOT scored by the current engine.',
    'work_arrangement': 'Collected/displayed; NOT scored by the current engine. Remote = no location penalty.',
    'type': 'LEGACY column (old vocabulary: Full-time/Part-time/Contract/Remote) — superseded by employment_type.',
    'work_setup': 'LEGACY column — superseded by work_arrangement.',
    'education_requirement': 'LEGACY column — superseded by min_education (still used as a fallback by normalize.ts).',
    'preferred_skills': 'LEGACY column — not read by the current engine.',
    'experience_requirement': 'LEGACY column — not read by the current engine.',
    'workplace_conditions': 'LEGACY column — folded into accommodations once, by normalize.ts, never read directly.',
    'accessibility_features': 'LEGACY column — folded into accommodations once, by normalize.ts, never read directly.',
    'physical_requirements': 'LEGACY column — not read by the current engine.',
    'communication_requirements': 'LEGACY column — not read by the current engine.',
    'functional_requirements': 'LEGACY column — not read by the current engine.',
    'accommodation_support': 'LEGACY column — folded into accommodations once, by normalize.ts, never read directly.',
    'match_percent': 'LEGACY column — not read by the current engine (scores are computed live, not stored).',
    'match_reasons': 'LEGACY column — not read by the current engine.',
    'screen_or_visual_demands': 'Added for accessibility matching; not currently read by the engine.',
    'accessibility_info': 'Free text shown on the job detail page; not scored.',
}
label_notes = {
    'pwd_user_id': 'Foreign key -> pwd_users.id.',
    'job_id': 'Foreign key -> jobs.id.',
    'saved': 'PROPOSED TARGET LABEL (0/1). No such column exists in the live database today — this simulates '
             'pwd_users.saved_job_ids (the bookmark button already in the app) as the closest available signal, '
             'since there is no explicit "liked" or "PDAO-marked match" column. See project chat for the full discussion.',
}

dd = (
    dd_rows_for(pwd_df, 'pwd_users', pwd_notes)
    + dd_rows_for(jobs_df, 'jobs', jobs_notes)
    + dd_rows_for(labels_df, 'recommendation_labels', label_notes)
)
dd_df = pd.DataFrame(dd)

# ── Write workbook ───────────────────────────────────────────────────────────
HEADER_FILL = PatternFill('solid', fgColor='0F766E')  # teal, matches the portal's brand color
HEADER_FONT = Font(bold=True, color='FFFFFF')

with pd.ExcelWriter('raw_dataset.xlsx', engine='openpyxl') as writer:
    pwd_df.to_excel(writer, sheet_name='pwd_users', index=False)
    jobs_df.to_excel(writer, sheet_name='jobs', index=False)
    labels_df.to_excel(writer, sheet_name='recommendation_labels', index=False)
    dd_df.to_excel(writer, sheet_name='Data_Dictionary', index=False)

    for name, df in [('pwd_users', pwd_df), ('jobs', jobs_df), ('recommendation_labels', labels_df), ('Data_Dictionary', dd_df)]:
        ws = writer.sheets[name]
        ws.freeze_panes = 'A2'
        for col_idx, col in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            width = min(45, max(12, df[col].astype(str).str.len().quantile(0.9) + 4))
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.row_dimensions[1].height = 26

print(f'pwd_users: {len(pwd_df)} rows')
print(f'jobs: {len(jobs_df)} rows')
print(f'recommendation_labels: {len(labels_df)} rows')
print(f'Data_Dictionary: {len(dd_df)} rows')
print('Wrote raw_dataset.xlsx')
